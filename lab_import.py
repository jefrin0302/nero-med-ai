"""
lab_import.py

Secure Laboratory Data & Clinical Assessment Import Parser for NeuroMed AI.
Supports CSV, JSON, PDF, and Word (.docx / .doc) formats based on the exact 23 clinical features.
Enforces:
1. Maximum upload size of 2 MB.
2. File format, extension, and document parse validation.
3. Strict column/key normalization and duplicate mapping detection.
4. Data type coercion, categorical normalization, and physiological boundary checks.
5. Strict preservation of missing values (NEVER substitutes zeros or fabricated numbers).
6. Separation of patient identity from the 23 clinical features.
7. Safe preview-only workflow: imported data is reviewed by clinician before submission.
"""

import csv
import io
import json
import re
from typing import Dict, Any, Tuple, Optional, List

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

MAX_UPLOAD_SIZE = 2 * 1024 * 1024  # 2 MB

# Canonical list of the exact 23 clinical features
CANONICAL_FEATURES = [
    "Age", "Sex", "Ceruloplasmin Level", "Copper in Blood Serum",
    "Free Copper in Blood Serum", "Copper in Urine", "ALT", "AST",
    "Total Bilirubin", "Albumin", "Alkaline Phosphatase (ALP)",
    "Prothrombin Time / INR", "Gamma-Glutamyl Transferase (GGT)",
    "Kayser-Fleischer Rings", "Neurological Symptoms Score",
    "Psychiatric Symptoms", "Cognitive Function Score",
    "Family History", "ATB7B Gene Mutation", "Region",
    "Socioeconomic Status", "Alcohol Use", "BMI"
]

# Map canonical feature names to HTML input IDs in templates/about.html
CANONICAL_TO_FORM_ID = {
    "Age": "inputAge",
    "Sex": "inputGender",
    "Region": "inputRegion",
    "Socioeconomic Status": "inputSocioeconomicStatus",
    "BMI": "inputBMI",
    "Alcohol Use": "inputAlcoholUse",
    "Ceruloplasmin Level": "inputCeruloplasmin",
    "Copper in Blood Serum": "inputCopperBlood",
    "Free Copper in Blood Serum": "inputFreeCopperBlood",
    "Copper in Urine": "inputCopperUrine",
    "ALT": "inputALT",
    "AST": "inputAST",
    "Total Bilirubin": "inputTotalBilirubin",
    "Albumin": "inputAlbumin",
    "Alkaline Phosphatase (ALP)": "inputALP",
    "Prothrombin Time / INR": "inputProthrombin",
    "Gamma-Glutamyl Transferase (GGT)": "inputGGT",
    "Kayser-Fleischer Rings": "inputKFR",
    "Neurological Symptoms Score": "inputNeurological",
    "Psychiatric Symptoms": "inputPsychiatric",
    "Cognitive Function Score": "inputCognitive",
    "Family History": "inputFamilyHistory",
    "ATB7B Gene Mutation": "inputGeneMutation"
}

# Recognized case-insensitive and snake_case aliases mapped to canonical keys
ALIAS_MAP = {
    # Demographics
    "age": "Age",
    "sex": "Sex",
    "gender": "Sex",
    "region": "Region",
    "geographic_region": "Region",
    "socioeconomic_status": "Socioeconomic Status",
    "ses": "Socioeconomic Status",
    "bmi": "BMI",
    "body_mass_index": "BMI",
    "alcohol_use": "Alcohol Use",
    "alcohol": "Alcohol Use",
    # Copper Panel
    "ceruloplasmin_level": "Ceruloplasmin Level",
    "ceruloplasmin": "Ceruloplasmin Level",
    "copper_in_blood_serum": "Copper in Blood Serum",
    "copper_blood_serum": "Copper in Blood Serum",
    "copper_blood": "Copper in Blood Serum",
    "serum_copper": "Copper in Blood Serum",
    "total_serum_copper": "Copper in Blood Serum",
    "free_copper_in_blood_serum": "Free Copper in Blood Serum",
    "free_copper_blood_serum": "Free Copper in Blood Serum",
    "free_copper": "Free Copper in Blood Serum",
    "free_serum_copper": "Free Copper in Blood Serum",
    "copper_in_urine": "Copper in Urine",
    "copper_urine": "Copper in Urine",
    "urine_copper": "Copper in Urine",
    "urinary_copper": "Copper in Urine",
    "24h_urinary_copper": "Copper in Urine",
    # LFTs & Enzymes
    "alt": "ALT",
    "alanine_aminotransferase": "ALT",
    "ast": "AST",
    "aspartate_aminotransferase": "AST",
    "total_bilirubin": "Total Bilirubin",
    "bilirubin": "Total Bilirubin",
    "albumin": "Albumin",
    "alkaline_phosphatase_alp": "Alkaline Phosphatase (ALP)",
    "alkaline_phosphatase": "Alkaline Phosphatase (ALP)",
    "alp": "Alkaline Phosphatase (ALP)",
    "prothrombin_time_inr": "Prothrombin Time / INR",
    "prothrombin_time": "Prothrombin Time / INR",
    "prothrombin": "Prothrombin Time / INR",
    "inr": "Prothrombin Time / INR",
    "gamma_glutamyl_transferase_ggt": "Gamma-Glutamyl Transferase (GGT)",
    "gamma_glutamyl_transferase": "Gamma-Glutamyl Transferase (GGT)",
    "ggt": "Gamma-Glutamyl Transferase (GGT)",
    # Neuro, Ophthalmic & Genetics
    "kayser_fleischer_rings": "Kayser-Fleischer Rings",
    "kayser_fleischer": "Kayser-Fleischer Rings",
    "kf_rings": "Kayser-Fleischer Rings",
    "kfr": "Kayser-Fleischer Rings",
    "neurological_symptoms_score": "Neurological Symptoms Score",
    "neurological_symptoms": "Neurological Symptoms Score",
    "neurological_score": "Neurological Symptoms Score",
    "neurological": "Neurological Symptoms Score",
    "psychiatric_symptoms": "Psychiatric Symptoms",
    "psychiatric": "Psychiatric Symptoms",
    "cognitive_function_score": "Cognitive Function Score",
    "cognitive_score": "Cognitive Function Score",
    "cognitive": "Cognitive Function Score",
    "family_history": "Family History",
    "atb7b_gene_mutation": "ATB7B Gene Mutation",
    "atp7b_gene_mutation": "ATB7B Gene Mutation",
    "atb7b_mutation": "ATB7B Gene Mutation",
    "atp7b_mutation": "ATB7B Gene Mutation",
    "atb7b": "ATB7B Gene Mutation",
    "atp7b": "ATB7B Gene Mutation",
    # Metadata Aliases
    "patient_name": "_patient_name",
    "name": "_patient_name",
    "patient_id": "_patient_id",
    "id": "_patient_id"
}

# Add exact canonical names in lower-case
for cf in CANONICAL_FEATURES:
    norm = re.sub(r'[^a-z0-9]', '_', cf.lower()).strip('_')
    ALIAS_MAP[norm] = cf
    ALIAS_MAP[cf.lower()] = cf

# Regex patterns for clinical feature extraction from unstructured document text
FEATURE_PATTERNS = {
    "Age": [
        r'(?:Patient\s+)?Age\s*[:=\t\s\n-]?\s*([0-9]{1,3})\b',
        r'\bAge\b.*?([0-9]{1,3})\s*(?:years|yrs|y\.o\.)'
    ],
    "Sex": [
        r'(?:Sex|Gender)\s*[:=\t\s\n-]?\s*(Male|Female|M|F)\b'
    ],
    "Region": [
        r'(?:Geographic\s+Region|Region)\s*[:=\t\s\n-]?\s*(East|North|South|West)\b'
    ],
    "Socioeconomic Status": [
        r'(?:Socioeconomic\s+Status|SES)\s*[:=\t\s\n-]?\s*(Low|Medium|High)\b'
    ],
    "BMI": [
        r'(?:Body\s+Mass\s+Index|BMI)\s*[:=\t\s\n()\-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Alcohol Use": [
        r'(?:Alcohol\s+Use|Alcohol)\s*[:=\t\s\n-]?\s*(True|False|Yes|No|Positive|Negative)\b'
    ],
    "Ceruloplasmin Level": [
        r'(?:Serum\s+)?Ceruloplasmin(?:\s+Level)?\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Copper in Blood Serum": [
        r'(?:Total\s+Serum\s+Copper|Copper\s+in\s+Blood(?:\s+Serum)?|Serum\s+Copper)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Free Copper in Blood Serum": [
        r'(?:Free\s+(?:\(Non-[\s\n]*Ceruloplasmin\)\s+)?Copper(?:\s+in\s+Blood(?:\s+Serum)?)?|Non-Ceruloplasmin\s+Copper)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Copper in Urine": [
        r'(?:24h?\s*(?:Urinary|Urine)?\s*Copper(?:\s*Excretion)?|Copper\s+in\s+Urine|Urine\s+Copper)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "ALT": [
        r'(?:Alanine\s+Aminotransferase(?:\s*\(ALT\))?|ALT)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "AST": [
        r'(?:Aspartate\s+Aminotransferase(?:\s*\(AST\))?|AST)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Total Bilirubin": [
        r'(?:Total\s+Bilirubin|Bilirubin)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Albumin": [
        r'(?:Serum\s+)?Albumin\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Alkaline Phosphatase (ALP)": [
        r'(?:Alkaline\s+Phosphatase(?:\s*\(ALP\))?|ALP)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Prothrombin Time / INR": [
        r'(?:Prothrombin\s+Time(?:\s*[\/\-]\s*INR)?|INR)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Gamma-Glutamyl Transferase (GGT)": [
        r'(?:Gamma-Glutamyl\s+Transferase(?:\s*\(GGT\))?|GGT)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)'
    ],
    "Kayser-Fleischer Rings": [
        r'(?:Kayser-Fleischer(?:\s*\(KF\))?\s*Rings?|KF\s*Rings?)\s*[:=\t\s\n-]*\s*([A-Za-z ()]+)'
    ],
    "Neurological Symptoms Score": [
        r'(?:Neurological\s+Symptoms(?:\s+Score)?)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)(?:\s*\/\s*10)?'
    ],
    "Psychiatric Symptoms": [
        r'(?:Documented\s+)?Psychiatric\s+Symptoms\s*[:=\t\s\n-]*\s*([A-Za-z ()]+)'
    ],
    "Cognitive Function Score": [
        r'(?:Cognitive\s+Function(?:\s+Score)?)\s*[:=\t\s\n-]*\s*([0-9]+(?:\.[0-9]+)?)(?:\s*\/\s*100)?'
    ],
    "Family History": [
        r'(?:Documented\s+)?Family\s+History(?:\s+of\s+Wilson(?:\'s)?\s+Disease)?\s*[:=\t\s\n-]*\s*([A-Za-z ()]+)'
    ],
    "ATB7B Gene Mutation": [
        r'(?:AT[PB]7B\s+Gene(?:\s+Mutation(?:\s+Analysis)?)?)\s*[:=\t\s\n-]*\s*(Mutation\s+(?:Present|Absent|Detected|Not\s+Detected)|Present|Absent|Positive|Negative|Yes|No|Wild\s*Type)\b',
        r'(?:AT[PB]7B\s+Gene(?:\s+Mutation(?:\s+Analysis)?)?)\s*[:=\t\s\n-]*\s*([A-Za-z ()]+)'
    ]
}


def normalize_column_name(raw_name: str) -> Optional[str]:
    """Resolves raw header/key to its canonical feature name or metadata alias."""
    if not raw_name or not isinstance(raw_name, str):
        return None
    cleaned = raw_name.strip()
    # Check exact match
    if cleaned in CANONICAL_FEATURES:
        return cleaned
    if cleaned in ("Patient Name", "Patient ID", "patient_name", "patient_id"):
        return "_patient_name" if "name" in cleaned.lower() else "_patient_id"

    # Normalize punctuation & underscores
    slug = re.sub(r'[^a-z0-9]', '_', cleaned.lower()).strip('_')
    return ALIAS_MAP.get(slug) or ALIAS_MAP.get(cleaned.lower())


def validate_and_coerce_field(canonical_name: str, raw_value: Any) -> Tuple[Optional[Any], Optional[str]]:
    """
    Validates data type, categoricals, and physiological bounds.
    Returns (coerced_value, error_message). If invalid, returns (None, reason).
    Preserves missing values as None.
    """
    if raw_value is None:
        return None, None
    s_val = str(raw_value).strip()
    if s_val == "" or s_val.lower() in ("null", "none", "nan", "n/a", "--", "-"):
        return None, None

    # Binary Fields
    binary_fields = {
        "Kayser-Fleischer Rings", "Psychiatric Symptoms",
        "Family History", "ATB7B Gene Mutation"
    }
    if canonical_name in binary_fields:
        s_low = s_val.lower()
        if any(term in s_low for term in ("0", "no", "false", "negative", "absent", "none", "wild type", "wild-type", "not observed", "not detected", "unremarkable")):
            return "No", None
        elif any(term in s_low for term in ("1", "yes", "true", "positive", "present", "observed", "detected", "pathogenic")):
            return "Yes", None
        return None, f"Expected Yes/No or 1/0, got '{raw_value}'"

    # Categorical: Sex
    if canonical_name == "Sex":
        s_low = s_val.lower()
        if s_low in ("male", "m") or s_low.startswith("male"):
            return "Male", None
        elif s_low in ("female", "f") or s_low.startswith("female"):
            return "Female", None
        return None, f"Expected Male or Female, got '{raw_value}'"

    # Categorical: Region
    if canonical_name == "Region":
        s_low = s_val.lower()
        valid_regions = {"east": "East", "north": "North", "south": "South", "west": "West"}
        for k, v in valid_regions.items():
            if k in s_low:
                return v, None
        return None, f"Expected East, North, South, or West, got '{raw_value}'"

    # Categorical: Socioeconomic Status
    if canonical_name == "Socioeconomic Status":
        s_low = s_val.lower()
        valid_ses = {"low": "Low", "medium": "Medium", "high": "High"}
        for k, v in valid_ses.items():
            if k in s_low:
                return v, None
        return None, f"Expected Low, Medium, or High, got '{raw_value}'"

    # Categorical: Alcohol Use
    if canonical_name == "Alcohol Use":
        s_low = s_val.lower()
        if any(term in s_low for term in ("0", "false", "no", "non-drinker", "none", "denies")):
            return "False", None
        elif any(term in s_low for term in ("1", "true", "yes", "positive", "drinker", "social", "moderate", "heavy")):
            return "True", None
        return None, f"Expected Yes/No or True/False, got '{raw_value}'"

    # Continuous Numerical Fields with Boundaries
    bounds = {
        "Age": (1.0, 120.0, "years (1–120)"),
        "BMI": (10.0, 65.0, "kg/m² (10–65)"),
        "Cognitive Function Score": (30.0, 100.0, "scale (30–100)"),
        "Ceruloplasmin Level": (0.0, 200.0, "mg/dL (0–200)"),
        "Copper in Blood Serum": (0.0, 500.0, "µg/dL (0–500)"),
        "Free Copper in Blood Serum": (0.0, 200.0, "µg/dL (0–200)"),
        "Copper in Urine": (0.0, 1000.0, "µg/24h (0–1000)"),
        "ALT": (0.0, 2000.0, "U/L (0–2000)"),
        "AST": (0.0, 2000.0, "U/L (0–2000)"),
        "Total Bilirubin": (0.0, 50.0, "mg/dL (0–50)"),
        "Albumin": (0.5, 8.0, "g/dL (0.5–8.0)"),
        "Alkaline Phosphatase (ALP)": (0.0, 1500.0, "U/L (0–1500)"),
        "Prothrombin Time / INR": (0.5, 10.0, "INR (0.5–10)"),
        "Gamma-Glutamyl Transferase (GGT)": (0.0, 1500.0, "U/L (0–1500)"),
        "Neurological Symptoms Score": (-10.0, 30.0, "score (-10–30)")
    }

    if canonical_name in bounds:
        min_v, max_v, desc = bounds[canonical_name]
        clean_num_str = re.sub(r'/(?:10|100)\b', '', s_val).strip()
        num_match = re.search(r'[-+]?[0-9]+(?:\.[0-9]+)?', clean_num_str)
        if not num_match:
            return None, f"Expected numerical value, got '{raw_value}'"
        try:
            num = float(num_match.group(0))
        except (ValueError, TypeError):
            return None, f"Expected numerical value, got '{raw_value}'"
        if num < min_v or num > max_v:
            return None, f"Value {num} outside clinical acceptable range {desc}"
        return num, None

    # Fallback numerical coercion
    try:
        return float(s_val), None
    except Exception:
        return s_val, None


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts raw text from PDF document using pypdf."""
    if not pypdf:
        raise RuntimeError("pypdf library is not installed on the system.")
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("PDF is encrypted / password-protected and cannot be read.")
        pages_text = []
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                pages_text.append(txt)
        return "\n".join(pages_text).strip()
    except Exception as e:
        if isinstance(e, ValueError):
            raise
        raise ValueError(f"Malformed or unreadable PDF document: {str(e)}")


def extract_content_from_docx(file_bytes: bytes) -> Tuple[str, List[List[str]]]:
    """Extracts paragraph text and structured table rows from Word (.docx) document."""
    if not docx:
        raise RuntimeError("python-docx library is not installed on the system.")
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        para_texts = [p.text for p in doc.paragraphs if p.text.strip()]
        table_rows: List[List[str]] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    table_rows.append(cells)
                    para_texts.append(" \t ".join(cells))
        return "\n".join(para_texts).strip(), table_rows
    except Exception as e:
        raise ValueError(f"Malformed or unreadable Word (.docx) document: {str(e)}")


def extract_text_from_doc(file_bytes: bytes) -> Tuple[str, List[List[str]]]:
    """Extracts text from legacy Word (.doc) binary OLE files or docx disguised as .doc."""
    if file_bytes.startswith(b"PK\x03\x04"):
        return extract_content_from_docx(file_bytes)

    extracted: List[str] = []
    # UTF-16LE extraction (typical for Microsoft Word binary OLE streams)
    raw_chars = []
    for i in range(0, len(file_bytes) - 1, 2):
        val = int.from_bytes(file_bytes[i:i+2], byteorder="little")
        if 32 <= val <= 126 or val in (9, 10, 13) or 160 <= val <= 255:
            raw_chars.append(chr(val))
        else:
            if len(raw_chars) >= 4:
                extracted.append("".join(raw_chars))
            raw_chars = []
    if len(raw_chars) >= 4:
        extracted.append("".join(raw_chars))

    # ASCII strings
    ascii_matches = re.findall(b'[\x20-\x7E\t\r\n]{4,}', file_bytes)
    for m in ascii_matches:
        try:
            extracted.append(m.decode("latin-1"))
        except Exception:
            pass

    full_txt = "\n".join(extracted).strip()
    return full_txt, []


def parse_unstructured_document(
    full_text: str,
    table_rows: Optional[List[List[str]]] = None
) -> Tuple[Optional[str], Optional[str], Dict[str, Any], Dict[str, Any], List[str], Dict[str, str], List[str]]:
    """
    Parses unstructured or semi-structured clinical document text and tables
    into canonical 23 features and patient metadata.
    Returns:
    (patient_name, patient_id, imported_fields, form_field_map, missing_fields, rejected_fields, unrecognized_cols)
    """
    patient_name: Optional[str] = None
    patient_id: Optional[str] = None
    candidates: Dict[str, Tuple[str, Any]] = {}
    rejected_fields: Dict[str, str] = {}
    imported_fields: Dict[str, Any] = {}
    form_field_map: Dict[str, Any] = {}
    unrecognized_cols: List[str] = []

    # 1. Extract patient metadata from text
    m_name = re.search(r'(?:Patient(?:\s+Full)?\s+Name|Full\s+Name|^Name)\s*[:=\t-]\s*([^\n\r,;|]+)', full_text, re.IGNORECASE | re.MULTILINE)
    if m_name:
        cand_name = m_name.group(1).strip()
        invalid_name_terms = ("measured value", "reference range", "clinical indicator", "parameter status", "specimen type", "sample type", "reference boundary")
        if len(cand_name) > 1 and not any(term in cand_name.lower() for term in invalid_name_terms):
            patient_name = cand_name

    m_id = re.search(r'(?:Patient\s+(?:ID|#|Identifier)|MRN|Record\s+(?:No|#)|Case\s+(?:ID|#))\s*[:=\t-]\s*([A-Za-z0-9_-]+)', full_text, re.IGNORECASE)
    if m_id:
        patient_id = m_id.group(1).strip()

    # 2. Extract from structured table rows if available (e.g. Word docx tables)
    for cells in (table_rows or []):
        if len(cells) < 2:
            continue
        raw_key = cells[0].strip()
        if not raw_key:
            continue
        canon = normalize_column_name(raw_key)
        if canon == "_patient_name" and not patient_name:
            patient_name = cells[1].strip()
        elif canon == "_patient_id" and not patient_id:
            patient_id = cells[1].strip()
        elif canon in CANONICAL_FEATURES:
            raw_val = cells[1].strip() if cells[1].strip() else (cells[2].strip() if len(cells) > 2 else "")
            if raw_val and canon not in candidates:
                candidates[canon] = (raw_key, raw_val)

    # 3. Pattern match across text for all canonical features
    for canon, patterns in FEATURE_PATTERNS.items():
        if canon in candidates:
            continue
        for pat in patterns:
            m = re.search(pat, full_text, re.IGNORECASE)
            if m:
                matched_val = m.group(1).strip()
                candidates[canon] = (canon, matched_val)
                break

    # 4. Validate and coerce all candidate values
    for canon, (source_col, raw_val) in candidates.items():
        coerced, err = validate_and_coerce_field(canon, raw_val)
        if err:
            rejected_fields[source_col] = err
        elif coerced is not None:
            imported_fields[canon] = coerced
            form_id = CANONICAL_TO_FORM_ID.get(canon)
            if form_id:
                form_field_map[form_id] = coerced

    missing_fields = [f for f in CANONICAL_FEATURES if f not in imported_fields]

    return patient_name, patient_id, imported_fields, form_field_map, missing_fields, rejected_fields, unrecognized_cols


def parse_lab_file_content(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses and validates CSV, JSON, PDF, or Word (.docx / .doc) laboratory file content.
    Returns:
    {
        "success": bool,
        "filename": str,
        "format": "csv" | "json" | "pdf" | "docx" | "doc",
        "patient_name": Optional[str],
        "patient_id": Optional[str],
        "imported_fields": Dict[str, Any],
        "form_field_map": Dict[str, Any],
        "missing_fields": List[str],
        "rejected_fields": Dict[str, str],
        "unrecognized_columns": List[str],
        "error": Optional[str]
    }
    """
    if len(file_bytes) > MAX_UPLOAD_SIZE:
        return {
            "success": False,
            "error": f"File size exceeds maximum permitted upload limit of {MAX_UPLOAD_SIZE // (1024*1024)} MB."
        }

    fname_lower = filename.lower()
    is_csv = fname_lower.endswith(".csv")
    is_json = fname_lower.endswith(".json")
    is_pdf = fname_lower.endswith(".pdf")
    is_docx = fname_lower.endswith(".docx")
    is_doc = fname_lower.endswith(".doc")

    if not any([is_csv, is_json, is_pdf, is_docx, is_doc]):
        return {
            "success": False,
            "error": "Unsupported file format. Please upload a valid laboratory report (.csv, .json, .pdf, .docx, or .doc)."
        }

    # Handle PDF
    if is_pdf:
        try:
            pdf_text = extract_text_from_pdf(file_bytes)
            if not pdf_text:
                return {
                    "success": False,
                    "error": "PDF file contains no extractable text or is empty. Please upload a digital PDF report."
                }
            p_name, p_id, imp_f, form_m, miss_f, rej_f, unrec_c = parse_unstructured_document(pdf_text)
            return {
                "success": True,
                "filename": filename,
                "format": "pdf",
                "patient_name": p_name,
                "patient_id": p_id,
                "imported_fields": imp_f,
                "form_field_map": form_m,
                "missing_fields": miss_f,
                "rejected_fields": rej_f,
                "unrecognized_columns": unrec_c,
                "total_canonical_count": len(CANONICAL_FEATURES),
                "imported_count": len(imp_f)
            }
        except Exception as e:
            return {"success": False, "error": f"Failed to parse PDF document: {str(e)}"}

    # Handle Word DOCX
    if is_docx:
        try:
            docx_text, table_rows = extract_content_from_docx(file_bytes)
            if not docx_text and not table_rows:
                return {
                    "success": False,
                    "error": "Word document (.docx) contains no readable text or tables."
                }
            p_name, p_id, imp_f, form_m, miss_f, rej_f, unrec_c = parse_unstructured_document(docx_text, table_rows)
            return {
                "success": True,
                "filename": filename,
                "format": "docx",
                "patient_name": p_name,
                "patient_id": p_id,
                "imported_fields": imp_f,
                "form_field_map": form_m,
                "missing_fields": miss_f,
                "rejected_fields": rej_f,
                "unrecognized_columns": unrec_c,
                "total_canonical_count": len(CANONICAL_FEATURES),
                "imported_count": len(imp_f)
            }
        except Exception as e:
            return {"success": False, "error": f"Failed to parse Word (.docx) document: {str(e)}"}

    # Handle Word DOC
    if is_doc:
        try:
            doc_text, table_rows = extract_text_from_doc(file_bytes)
            if not doc_text:
                return {
                    "success": False,
                    "error": "Legacy Word document (.doc) contains no readable text or is corrupted."
                }
            p_name, p_id, imp_f, form_m, miss_f, rej_f, unrec_c = parse_unstructured_document(doc_text, table_rows)
            return {
                "success": True,
                "filename": filename,
                "format": "doc",
                "patient_name": p_name,
                "patient_id": p_id,
                "imported_fields": imp_f,
                "form_field_map": form_m,
                "missing_fields": miss_f,
                "rejected_fields": rej_f,
                "unrecognized_columns": unrec_c,
                "total_canonical_count": len(CANONICAL_FEATURES),
                "imported_count": len(imp_f)
            }
        except Exception as e:
            return {"success": False, "error": f"Failed to parse Word (.doc) document: {str(e)}"}

    # Handle JSON and CSV
    raw_dict: Dict[str, Any] = {}

    if is_json:
        try:
            text = file_bytes.decode("utf-8-sig", errors="replace")
            parsed = json.loads(text)
            if isinstance(parsed, list):
                if len(parsed) == 0 or not isinstance(parsed[0], dict):
                    return {"success": False, "error": "JSON list is empty or does not contain patient record objects."}
                raw_dict = parsed[0]
            elif isinstance(parsed, dict):
                raw_dict = parsed
            else:
                return {"success": False, "error": "JSON root must be an object or a list of record objects."}
        except Exception as e:
            return {"success": False, "error": f"Malformed JSON file: {str(e)}"}

    if is_csv:
        try:
            text = file_bytes.decode("utf-8-sig", errors="replace")
            sample_line = text.splitlines()[0] if text.splitlines() else ""
            delimiter = "\t" if "\t" in sample_line and "," not in sample_line else ","
            raw_reader = list(csv.reader(io.StringIO(text), delimiter=delimiter))
            if not raw_reader:
                return {"success": False, "error": "CSV file is completely empty."}
            headers = [h.strip() for h in raw_reader[0] if h is not None]
            if len(raw_reader) < 2:
                return {"success": False, "error": "CSV file contains headers but no patient data rows."}
            
            seen_headers = {}
            for h in headers:
                if not h:
                    continue
                h_low = h.lower()
                seen_headers[h_low] = seen_headers.get(h_low, 0) + 1
            duplicate_headers = [h for h, count in seen_headers.items() if count > 1]
            if duplicate_headers:
                return {
                    "success": False,
                    "error": f"CSV contains duplicate column headers: {', '.join(duplicate_headers)}. Please remove redundant columns."
                }

            reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
            rows = list(reader)
            if not rows:
                return {"success": False, "error": "CSV file contains no data rows."}
            raw_dict = rows[0]
        except Exception as e:
            return {"success": False, "error": f"Malformed CSV file: {str(e)}"}

    patient_name = None
    patient_id = None
    canonical_candidates: Dict[str, List[Tuple[str, Any]]] = {}
    unrecognized_cols: List[str] = []

    for raw_k, raw_v in raw_dict.items():
        if raw_k is None:
            continue
        k_str = str(raw_k).strip()
        if k_str.startswith(("_comment", "_notice", "#")):
            continue
        canon = normalize_column_name(k_str)
        if canon == "_patient_name":
            if raw_v is not None and str(raw_v).strip():
                patient_name = str(raw_v).strip()
        elif canon == "_patient_id":
            if raw_v is not None and str(raw_v).strip():
                patient_id = str(raw_v).strip()
        elif canon in CANONICAL_FEATURES:
            canonical_candidates.setdefault(canon, []).append((k_str, raw_v))
        else:
            unrecognized_cols.append(k_str)

    rejected_fields: Dict[str, str] = {}
    imported_fields: Dict[str, Any] = {}
    form_field_map: Dict[str, Any] = {}

    for canon, candidates in canonical_candidates.items():
        if len(candidates) > 1:
            unique_vals = set(str(v).strip() for _, v in candidates if v is not None and str(v).strip())
            if len(unique_vals) > 1:
                col_names = ", ".join(f"'{c}'" for c, _ in candidates)
                rejected_fields[canon] = f"Conflicting duplicate columns detected for {canon}: {col_names}"
                continue
        col_used, val_raw = candidates[0]
        coerced_val, err = validate_and_coerce_field(canon, val_raw)
        if err:
            rejected_fields[col_used] = err
        elif coerced_val is not None:
            imported_fields[canon] = coerced_val
            form_id = CANONICAL_TO_FORM_ID.get(canon)
            if form_id:
                form_field_map[form_id] = coerced_val

    missing_fields = [f for f in CANONICAL_FEATURES if f not in imported_fields]

    return {
        "success": True,
        "filename": filename,
        "format": "csv" if is_csv else "json",
        "patient_name": patient_name,
        "patient_id": patient_id,
        "imported_fields": imported_fields,
        "form_field_map": form_field_map,
        "missing_fields": missing_fields,
        "rejected_fields": rejected_fields,
        "unrecognized_columns": unrecognized_cols,
        "total_canonical_count": len(CANONICAL_FEATURES),
        "imported_count": len(imported_fields)
    }


def generate_sample_csv_text() -> str:
    """Generates downloadable sample CSV template text matching the exact 23 features."""
    headers = [
        "Patient ID", "Patient Name", "Age", "Sex", "Region", "Socioeconomic Status",
        "BMI", "Alcohol Use", "Ceruloplasmin Level", "Copper in Blood Serum",
        "Free Copper in Blood Serum", "Copper in Urine", "ALT", "AST", "Total Bilirubin",
        "Albumin", "Alkaline Phosphatase (ALP)", "Prothrombin Time / INR",
        "Gamma-Glutamyl Transferase (GGT)", "Kayser-Fleischer Rings",
        "Neurological Symptoms Score", "Psychiatric Symptoms", "Cognitive Function Score",
        "Family History", "ATB7B Gene Mutation"
    ]
    sample_row = [
        "NEUROMED001", "Sample Patient (Fictional)", "24", "Female", "North", "High",
        "22.5", "False", "28.5", "115.0",
        "8.2", "24.5", "22.0", "25.0", "0.80",
        "4.4", "85.0", "1.05",
        "28.0", "No",
        "0.2", "No", "95.0",
        "No", "No"
    ]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerow(sample_row)
    return output.getvalue()


def generate_sample_json_text() -> str:
    """Generates downloadable sample JSON template matching the exact 23 features."""
    sample_data = {
        "_notice": "Fictional sample data for demonstration only. Not a real patient.",
        "Patient ID": "NEUROMED001",
        "Patient Name": "Sample Patient (Fictional)",
        "Age": 24,
        "Sex": "Female",
        "Region": "North",
        "Socioeconomic Status": "High",
        "BMI": 22.5,
        "Alcohol Use": False,
        "Ceruloplasmin Level": 28.5,
        "Copper in Blood Serum": 115.0,
        "Free Copper in Blood Serum": 8.2,
        "Copper in Urine": 24.5,
        "ALT": 22.0,
        "AST": 25.0,
        "Total Bilirubin": 0.80,
        "Albumin": 4.4,
        "Alkaline Phosphatase (ALP)": 85.0,
        "Prothrombin Time / INR": 1.05,
        "Gamma-Glutamyl Transferase (GGT)": 28.0,
        "Kayser-Fleischer Rings": "No",
        "Neurological Symptoms Score": 0.2,
        "Psychiatric Symptoms": "No",
        "Cognitive Function Score": 95.0,
        "Family History": "No",
        "ATB7B Gene Mutation": "No"
    }
    return json.dumps(sample_data, indent=2)


def generate_sample_docx_bytes() -> bytes:
    """Generates downloadable sample Word (.docx) report matching the exact 23 features."""
    if not docx:
        raise RuntimeError("python-docx is not installed on the system.")
    doc = docx.Document()
    doc.add_heading("Wilson's Disease Laboratory Report (Sample)", 0)
    p_meta = doc.add_paragraph()
    p_meta.add_run("Patient Name: ").bold = True
    p_meta.add_run("Sample Patient (Fictional)\n")
    p_meta.add_run("Patient ID: ").bold = True
    p_meta.add_run("NEUROMED001\n")
    p_meta.add_run("Specimen Type: ").bold = True
    p_meta.add_run("Serum, 24h Urine, Whole Blood EDTA")

    doc.add_heading("Clinical Indicators & Biochemistry Panel", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Clinical Biomarker / Parameter"
    hdr_cells[1].text = "Measured Value"
    hdr_cells[2].text = "Standard Reference / Unit"

    rows_data = [
        ("Age", "24", "years (1–120)"),
        ("Sex", "Female", "Male / Female"),
        ("Region", "North", "East / North / South / West"),
        ("Socioeconomic Status", "High", "Low / Medium / High"),
        ("BMI", "22.5", "kg/m² (10–65)"),
        ("Alcohol Use", "False", "True / False"),
        ("Ceruloplasmin Level", "28.5", "mg/dL (20–40 mg/dL)"),
        ("Copper in Blood Serum", "115.0", "µg/dL (70–140 µg/dL)"),
        ("Free Copper in Blood Serum", "8.2", "µg/dL (<15 µg/dL)"),
        ("Copper in Urine", "24.5", "µg/24h (<40 µg/24h)"),
        ("ALT", "22.0", "U/L (<40 U/L)"),
        ("AST", "25.0", "U/L (<40 U/L)"),
        ("Total Bilirubin", "0.80", "mg/dL (0.2–1.2 mg/dL)"),
        ("Albumin", "4.4", "g/dL (3.5–5.0 g/dL)"),
        ("Alkaline Phosphatase (ALP)", "85.0", "U/L (40–130 U/L)"),
        ("Prothrombin Time / INR", "1.05", "INR (0.8–1.2)"),
        ("Gamma-Glutamyl Transferase (GGT)", "28.0", "U/L (<50 U/L)"),
        ("Kayser-Fleischer Rings", "Absent", "Negative (Absent)"),
        ("Neurological Symptoms Score", "0.2", "Clinical Scale (0–10)"),
        ("Psychiatric Symptoms", "Absent", "None Documented"),
        ("Cognitive Function Score", "95.0", "Scale (30–100)"),
        ("Family History", "Negative", "No Known Family History"),
        ("ATB7B Gene Mutation", "Absent", "Wild Type (No Mutation)")
    ]

    for item, val, ref in rows_data:
        row = table.add_row().cells
        row[0].text = item
        row[1].text = val
        row[2].text = ref

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
